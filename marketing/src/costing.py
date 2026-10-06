from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter

OUT = "/home/user/InteractiveBirthdayInvitation/marketing/digital-invitation-costing-PH.xlsx"
F = "Arial"
BLUE = Font(name=F, color="0000FF"); BLACK = Font(name=F); GREEN = Font(name=F, color="008000")
BOLD = Font(name=F, bold=True); TITLE = Font(name=F, bold=True, size=14, color="5C3A63")
HDR = Font(name=F, bold=True, color="FFFFFF"); NOTE = Font(name=F, italic=True, color="666666", size=9)
YELLOW = PatternFill("solid", fgColor="FFFF00"); HFILL = PatternFill("solid", fgColor="5C3A63")
SUB = PatternFill("solid", fgColor="EFE3F2"); TOT = PatternFill("solid", fgColor="F6EEF8")
thin = Side(style="thin", color="BBBBBB"); TOP = Border(top=Side(style="thin", color="5C3A63"))
PESO = '"₱"#,##0;("₱"#,##0);"-"'; USD = '"US$"#,##0.00'; PCT = '0.0%;(0.0%);"-"'; HRS = '0.0" h"'

wb = Workbook()

def sheet(name, title, widths):
    ws = wb.create_sheet(name)
    ws["A1"] = title; ws["A1"].font = TITLE
    for i, w in enumerate(widths): ws.column_dimensions[get_column_letter(i + 1)].width = w
    ws.sheet_view.showGridLines = False
    return ws

def header(ws, row, labels):
    for i, l in enumerate(labels):
        c = ws.cell(row, i + 1, l); c.font = HDR; c.fill = HFILL
        c.alignment = Alignment(horizontal="center" if i else "left", vertical="center", wrap_text=True)

def put(ws, ref, v, font=BLACK, fmt=None, fill=None, bold=False):
    c = ws[ref]; c.value = v
    c.font = Font(name=F, bold=bold, color=font.color) if bold else font
    if fmt: c.number_format = fmt
    if fill: c.fill = fill
    return c

# ---------------- Inputs ----------------
inp = sheet("Inputs", "Inputs & assumptions (edit the blue cells)", [44, 14, 16, 70])
inp["A2"] = "Blue = your inputs, yellow = the ones worth reviewing first. Everything else in the workbook recalculates from here."; inp["A2"].font = NOTE
header(inp, 3, ["Assumption", "Value", "Unit", "Source / note"])
rows = [
    (4, "Exchange rate", 61, "₱ per US$", "Rate on 3 Aug 2026 was ₱60.92 (valutafx.com). Update to today's rate.", '"₱"0.00', True),
    (5, "Your labour rate", 250, "₱ per hour", "Assumption: entry-level PH freelance web rate. Raise as you get faster/busier.", PESO, True),
    (6, "Vercel Pro", 20, "US$ per month", "Hosting for all sites + email function. Hobby (free) plan does not allow commercial use, so Pro is required (vercel.com pricing; makerkit.dev/blog/saas/vercel-cost).", USD, False),
    (7, "Supabase Pro", 25, "US$ per month", "Database + photo storage. Free tier: 2 projects, 1 GB storage, pauses after 1 week idle (makerkit.dev/blog/md/saas/supabase-pricing).", USD, False),
    (8, "Use Supabase Pro? (1 = yes, 0 = free tier)", 1, "switch", "Start on the free tier (0) for your first few clients; switch to Pro (1) once photos pass ~1 GB or a site risks pausing. Assumes all events share one project (each event gets its own table and photo folder; the SQL files need the table name changed per event).", "0", True),
    (9, "Custom .com domain", 12, "US$ per year", "Assumption: typical .com registration (Namecheap/Porkbun range). Only charged on Signature or as an add-on.", USD, False),
    (10, "Design tools (e.g. Canva Pro)", 0, "₱ per month", "Optional. Set your actual subscription if you use one.", PESO, False),
    (11, "Instagram / Facebook ads", 2400, "₱ per month", "Assumption: ₱80/day boosted story or reel. Set 0 if you only post organically.", PESO, True),
    (12, "Payment fee", 0.02, "% of price", "Assumption for GCash/Maya business QR or card. Use 0% for personal GCash/bank transfer.", PCT, False),
    (13, "Income tax (8% option)", 0.08, "% of price", "BIR 8% optional rate on gross sales for self-employed (first ₱250,000/yr is exempt, so this is conservative). Confirm with BIR/an accountant; register your business first.", PCT, False),
    (14, "Gmail for confirmation emails", 0, "₱ per month", "Free Gmail app password; ~500 emails/day limit is plenty per event.", PESO, False),
]
for r, label, val, unit, note, fmt, key in rows:
    inp.cell(r, 1, label).font = BLACK
    put(inp, f"B{r}", val, BLUE, fmt, YELLOW if key else None)
    inp.cell(r, 3, unit).font = BLACK
    c = inp.cell(r, 4, note); c.font = NOTE; c.alignment = Alignment(wrap_text=True, vertical="top")
header(inp, 16, ["Expected events per month (package mix)", "Events", "", "Note"])
for r, (n, v) in zip((17, 18, 19), [("Basic", 3), ("Classic", 4), ("Signature", 1)]):
    inp.cell(r, 1, n).font = BLACK; put(inp, f"B{r}", v, BLUE, "0", YELLOW)
inp["D17"] = "Assumption: a realistic early month. Change it to see profit at your real volume."; inp["D17"].font = NOTE
put(inp, "A20", "Total events per month", bold=True); put(inp, "B20", "=SUM(B17:B19)", BLACK, "0", bold=True)
for r in range(4, 21): inp.row_dimensions[r].height = 30 if r < 15 else 18

# ---------------- Monthly Costs ----------------
mc = sheet("Monthly Costs", "Monthly running costs (₱)", [42, 16, 60])
header(mc, 3, ["Cost", "₱ / month", "How it's calculated"])
lines = [
    (4, "Vercel Pro", "=Inputs!B6*Inputs!B4", "US$ price × exchange rate"),
    (5, "Supabase Pro", "=Inputs!B7*Inputs!B4*Inputs!B8", "US$ price × exchange rate × Pro switch"),
    (6, "Gmail", "=Inputs!B14", "Free"),
    (7, "Design tools", "=Inputs!B10", "From Inputs"),
]
for r, l, f, how in lines:
    mc.cell(r, 1, l).font = BLACK; put(mc, f"B{r}", f, GREEN, PESO); mc.cell(r, 3, how).font = NOTE
put(mc, "A8", "Platform & tools subtotal", bold=True); put(mc, "B8", "=SUM(B4:B7)", BLACK, PESO, bold=True)
mc["A9"] = "Instagram / Facebook ads"; mc["A9"].font = BLACK; put(mc, "B9", "=Inputs!B11", GREEN, PESO)
for col in "AB": mc[f"{col}10"].fill = TOT; mc[f"{col}10"].border = TOP
put(mc, "A10", "Total monthly overhead", bold=True); put(mc, "B10", "=B8+B9", BLACK, PESO, TOT, bold=True)
mc["A12"] = "Events per month"; mc["A12"].font = BLACK; put(mc, "B12", "=Inputs!B20", GREEN, "0")
mc["A13"] = "Platform & tools per event"; mc["A13"].font = BLACK; put(mc, "B13", "=IFERROR(B8/B12,0)", BLACK, PESO)
mc["C13"] = "Overhead spread across the month's events"; mc["C13"].font = NOTE
mc["A14"] = "Ads per event"; mc["A14"].font = BLACK; put(mc, "B14", "=IFERROR(B9/B12,0)", BLACK, PESO)

# ---------------- Packages ----------------
pk = wb["Sheet"]; pk.title = "Packages"; wb.move_sheet(pk, -10)
pk["A1"] = "Digital invitation packages: price & cost per event (Philippines)"; pk["A1"].font = TITLE
pk["A2"] = "Prices in blue are yours to set. Costs pull from Inputs and Monthly Costs."; pk["A2"].font = NOTE
pk.sheet_view.showGridLines = False
for col, w in zip("ABCDE", [40, 17, 17, 17, 54]): pk.column_dimensions[col].width = w
header(pk, 4, ["", "Basic", "Classic", "Signature", "Note"])
put(pk, "A5", "Best for", bold=True)
for col, v in zip("BCD", ["Simple parties, kids' birthdays", "Debuts, 18ths, sweet 16s (most popular)", "Weddings, big debuts"]):
    c = pk[f"{col}5"]; c.value = v; c.font = NOTE; c.alignment = Alignment(wrap_text=True, horizontal="center")
pk.row_dimensions[5].height = 30
put(pk, "A6", "Package price (₱)", bold=True)
for col, v in zip("BCD", [1999, 3499, 5499]): put(pk, f"{col}6", v, BLUE, PESO, YELLOW, bold=True)
pk["E6"] = "Positioned against PH market (₱800–₱5,997, see Market tab); more features than most at each tier."; pk["E6"].font = NOTE
pk["A7"] = "Includes custom domain (1 = yes)"; pk["A7"].font = BLACK
for col, v in zip("BCD", [0, 0, 1]): put(pk, f"{col}7", v, BLUE, "0")

put(pk, "A9", "Your time per event (hours)", bold=True)
for col in "ABCDE": pk[f"{col}9"].fill = SUB
hours = [("Setup & configuration", [0.5, 1, 1]), ("Design & customisation", [1, 3, 4.5]), ("Revisions", [0.5, 1.5, 2.5]),
         ("Client chat & guest support", [0.5, 1, 1.5]), ("Testing & launch", [0.5, 0.5, 0.5])]
for i, (l, vs) in enumerate(hours):
    r = 10 + i; pk.cell(r, 1, l).font = BLACK
    for col, v in zip("BCD", vs): put(pk, f"{col}{r}", v, BLUE, HRS)
pk["E10"] = "Assumption: the template + new_event.py script does most setup; adjust after your first few clients."; pk["E10"].font = NOTE
put(pk, "A15", "Total hours", bold=True)
for col in "BCD": put(pk, f"{col}15", f"=SUM({col}10:{col}14)", BLACK, HRS, bold=True)

put(pk, "A17", "Cost per event (₱)", bold=True)
for col in "ABCDE": pk[f"{col}17"].fill = SUB
cost_rows = [
    (18, "Your labour", "={c}15*Inputs!$B$5", GREEN, "Hours × your labour rate"),
    (19, "Platform & tools share", "='Monthly Costs'!$B$13", GREEN, "Vercel + Supabase + tools ÷ events per month"),
    (20, "Ads share", "='Monthly Costs'!$B$14", GREEN, "Monthly ads ÷ events per month"),
    (21, "Custom domain", "={c}7*Inputs!$B$9*Inputs!$B$4", GREEN, "Signature only (1 year)"),
    (22, "Payment fee", "={c}6*Inputs!$B$12", GREEN, "Price × payment fee %"),
    (23, "Income tax", "={c}6*Inputs!$B$13", GREEN, "Price × 8% (conservative)"),
]
for r, l, f, font, note in cost_rows:
    pk.cell(r, 1, l).font = BLACK; pk.cell(r, 5, note).font = NOTE
    for col in "BCD": put(pk, f"{col}{r}", f.format(c=col), font, PESO)
put(pk, "A24", "Total cost per event", bold=True)
for col in "BCD":
    put(pk, f"{col}24", f"=SUM({col}18:{col}23)", BLACK, PESO, bold=True); pk[f"{col}24"].border = TOP

put(pk, "A26", "Profit (₱)", bold=True)
for col in "ABCDE": pk[f"{col}26"].fill = SUB
prof = [
    (27, "Business profit (after paying yourself)", "={c}6-{c}24", PESO, "What the business keeps after your hourly pay"),
    (28, "Net margin", "=IFERROR({c}27/{c}6,0)", PCT, ""),
    (29, "Your take-home (profit + labour)", "={c}27+{c}18", PESO, "If you do the work yourself, this is what you pocket"),
    (30, "Effective earnings per hour", "=IFERROR({c}29/{c}15,0)", PESO, "Take-home ÷ hours"),
    (31, "Cash cost (excluding your time)", "={c}24-{c}18", PESO, "Money that actually leaves your pocket"),
]
for r, l, f, fmt, note in prof:
    pk.cell(r, 1, l).font = BOLD if r in (27, 29) else BLACK; pk.cell(r, 5, note).font = NOTE
    for col in "BCD":
        c = put(pk, f"{col}{r}", f.format(c=col), BLACK, fmt, TOT if r in (27, 29) else None, bold=r in (27, 29))
    if r in (27, 29): pk[f"A{r}"].fill = TOT

put(pk, "A33", "What's included", bold=True)
for col in "ABCDE": pk[f"{col}33"].fill = SUB
feats = [
    ("Envelope that opens to the invitation card", "✓", "✓", "✓"),
    ("Live countdown + add-to-calendar (Google, Apple, Outlook)", "✓", "✓", "✓"),
    ("Map & one-tap directions", "✓", "✓", "✓"),
    ("Online RSVP (names, emails, wishes)", "✓", "✓", "✓"),
    ("Private guest list for hosts (see who's coming)", "✓", "✓", "✓"),
    ("Email confirmations with calendar invite", "—", "✓", "✓"),
    ("Guest photo wall + all-photos page", "—", "✓", "✓"),
    ("Theme: colours, fonts, wording", "Preset", "Custom", "Custom"),
    ("Decorations / motifs redrawn for your theme", "—", "—", "✓"),
    ("Invitation card design", "Your own card", "Your own card", "Designed for you"),
    ("Web address", "yourname.vercel.app", "yourname.vercel.app", "Custom .com (1 yr)"),
    ("Revision rounds", "1", "2", "3"),
    ("Turnaround", "3–5 days", "3–5 days", "5–7 days"),
    ("Site stays online", "1 month after event", "3 months after event", "12 months after event"),
]
for i, row in enumerate(feats):
    r = 34 + i; pk.cell(r, 1, row[0]).font = BLACK
    for col, v in zip("BCD", row[1:]):
        c = pk[f"{col}{r}"]; c.value = v; c.font = BLACK; c.alignment = Alignment(horizontal="center", wrap_text=True)
pk.freeze_panes = "B5"

# ---------------- Add-ons ----------------
ad = sheet("Add-ons", "Add-ons (₱)", [44, 14, 14, 14, 62])
header(ad, 3, ["Add-on", "Price", "Your cost", "Profit", "Cost basis"])
addons = [
    ("Rush delivery (48 hours)", 800, "=1*Inputs!$B$5", "~1 hour of overtime/reshuffling; you are mostly charging for priority"),
    ("Extra revision round", 400, "=1*Inputs!$B$5", "1 hour"),
    ("Custom .com domain (1 year)", 1000, "=Inputs!$B$9*Inputs!$B$4", "Domain registration"),
    ("Invitation card design (if you don't have one)", 1200, "=3*Inputs!$B$5", "3 hours"),
    ("Photo wall (Basic package)", 700, "=1*Inputs!$B$5", "1 hour setup"),
    ("Email confirmations (Basic package)", 500, "=1*Inputs!$B$5", "1 hour setup"),
    ("Keep site online +6 months", 500, "=0.5*Inputs!$B$5", "30 min upkeep; hosting already covered"),
    ("Download all guest photos (ZIP)", 300, "=0.5*Inputs!$B$5", "30 min"),
    ("QR code for printed invites", 150, "=0.25*Inputs!$B$5", "15 min"),
    ("Instagram story teaser video of their invite", 800, "=2*Inputs!$B$5", "2 hours (reuse the promo video template)"),
]
for i, (n, p, cost, basis) in enumerate(addons):
    r = 4 + i; ad.cell(r, 1, n).font = BLACK
    put(ad, f"B{r}", p, BLUE, PESO); put(ad, f"C{r}", cost, GREEN, PESO); put(ad, f"D{r}", f"=B{r}-C{r}", BLACK, PESO)
    ad.cell(r, 5, basis).font = NOTE
ad["A15"] = "Prices are suggestions (assumptions); hours are valued at your labour rate on Inputs."; ad["A15"].font = NOTE

# ---------------- Break-even ----------------
be = sheet("Break-even", "Break-even & monthly projection", [44, 17, 17, 17, 17])
put(be, "A3", "Monthly overhead (platform + ads)"); put(be, "B3", "='Monthly Costs'!B10", GREEN, PESO, bold=True)
header(be, 5, ["Per event", "Basic", "Classic", "Signature", "Total / mix"])
be_rows = [
    (6, "Price", "=Packages!{c}6"),
    (7, "Variable cost (labour, domain, fees, tax)", "=Packages!{c}18+Packages!{c}21+Packages!{c}22+Packages!{c}23"),
    (8, "Contribution to overhead", "={c}6-{c}7"),
]
for r, l, f in be_rows:
    be.cell(r, 1, l).font = BLACK
    for col in "BCD": put(be, f"{col}{r}", f.format(c=col), GREEN if r < 8 else BLACK, PESO)
be["A9"] = "Events needed to cover overhead (one package only)"; be["A9"].font = BOLD
for col in "BCD": put(be, f"{col}9", f"=IFERROR(ROUNDUP($B$3/{col}8,0),0)", BLACK, "0", bold=True)

header(be, 11, ["This month (from Inputs mix)", "Basic", "Classic", "Signature", "Total"])
proj = [
    (12, "Events", "=Inputs!B{m}", "0", GREEN),
    (13, "Revenue", "={c}6*{c}12", PESO, BLACK),
    (14, "Variable costs", "={c}7*{c}12", PESO, BLACK),
    (15, "Contribution", "={c}13-{c}14", PESO, BLACK),
]
for r, l, f, fmt, font in proj:
    be.cell(r, 1, l).font = BLACK
    for col, m in zip("BCD", (17, 18, 19)): put(be, f"{col}{r}", f.format(c=col, m=m), font, fmt)
    put(be, f"E{r}", f"=SUM(B{r}:D{r})", BLACK, fmt, bold=True)
be["A16"] = "Less: monthly overhead"; be["A16"].font = BLACK; put(be, "E16", "=-B3", BLACK, PESO)
put(be, "A17", "Business profit this month", bold=True); put(be, "E17", "=E15+E16", BLACK, PESO, TOT, bold=True); be["A17"].fill = TOT
be["A18"] = "Plus your labour pay"; be["A18"].font = BLACK
put(be, "E18", "=SUMPRODUCT(Packages!B18:D18,B12:D12)", GREEN, PESO)
put(be, "A19", "Your total take-home this month", bold=True); put(be, "E19", "=E17+E18", BLACK, PESO, TOT, bold=True); be["A19"].fill = TOT
be["A20"] = "Total hours worked"; be["A20"].font = BLACK
put(be, "E20", "=SUMPRODUCT(Packages!B15:D15,B12:D12)", GREEN, HRS)

# ---------------- Market ----------------
mk = sheet("Market", "Philippine market check (Oct 2026)", [30, 24, 14, 50, 70])
header(mk, 3, ["Provider", "Package", "Price", "What's mentioned", "Source"])
market = [
    ("E-Invitation (Logica Technology)", "Standard", "₱800", "RSVP, mobile sharing, fast delivery", "https://e-invitation.online/"),
    ("Click-Invitation", "Starter Package 1", "₱1,998", "Wedding digital invitation with RSVP", "https://click-invitation.com/blog/affordable-wedding-digital-invitations-in-the-philippines-starter-package-1-for-only-php1-998"),
    ("Click-Invitation", "Starter Package 2", "₱2,998", "RSVP tracking, online guest management", "https://click-invitation.com/blog/starter-package-2-the-complete-wedding-digital-invitation-experience-for-only-php2-998"),
    ("Specially Invited", "Essential", "from ₱2,497", "Sub-domain website invitation, live RSVP, 2–5 days", "https://www.speciallyinvited.com/"),
    ("Specially Invited", "Signature", "₱3,997", "Personalised domain option", "https://www.speciallyinvited.com/"),
    ("Specially Invited", "Ultimate", "₱5,997", "Bespoke, domain-ready", "https://www.speciallyinvited.com/"),
    ("Etsy templates", "DIY template", "≈US$3–33", "Template only; you edit and host yourself", "https://www.etsy.com/market/filipino_digital_invitation"),
]
for i, row in enumerate(market):
    for j, v in enumerate(row):
        c = mk.cell(4 + i, j + 1, v); c.font = NOTE if j == 4 else BLACK
mk["A12"] = "Prices as listed in web search results in Oct 2026; check the sites for current offers. Our photo wall, email confirmations and host guest list are rarely included at the lower tiers."; mk["A12"].font = NOTE

wb.save(OUT)
print(OUT)
